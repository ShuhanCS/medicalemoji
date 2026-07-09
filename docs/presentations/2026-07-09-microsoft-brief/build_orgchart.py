"""Replace slide 7 of the brief with a Unicode org chart showing where Microsoft sits
relative to the emoji working group.

Reads Shuhan's edited 8-slide deck and rewrites only that one slide.
"""
from pptx.enum.text import PP_ALIGN

from deck_kit import AMBER, BLUE, HEAD, INK, MID, SLATE, WHITE, Deck, lines, run

SRC = r"C:\Users\Shuha\Downloads\Emoji-2026-Brief-v2.pptx"
OUT = "Emoji-2026-Brief-v3.pptx"
TITLE = "Inside the Unicode Consortium"
S = "org"

deck = Deck(SRC)
old = deck.find_slide("Where emoji decisions get made")
# Add before dropping: a new part created after a drop would reuse the freed
# partname (slide8.xml) and silently overwrite the Request slide.
s = deck.base(TITLE, S)
deck.drop_slide(old)
deck.move_slide(s, 6)  # 0-indexed: becomes slide 7, ahead of "Request"

tf = deck.txbox(s, S, "sub", 0.62, 1.06, 8.66, 0.30)
p = tf.paragraphs[0]
run(p, "Microsoft sits on the board and votes in the UTC. ", size=10.5, color=SLATE)
run(p, "Emoji sits one level below the UTC.", size=10.5, bold=True, color=INK)

# ---------------------------------------------------------------- geometry
COLW, GAP, X0 = 1.66, 0.12, 0.61
CENTERS = [X0 + i * (COLW + GAP) + COLW / 2 for i in range(5)]
MIDX = CENTERS[2]

ROOT_Y, ROOT_H, ROOT_W = 1.44, 0.40, 2.20
BUS1_Y = 2.02
ROW1_Y, ROW1_H = 2.16, 0.88
BUS2_Y = 3.18
ROW2_Y, ROW2_H = 3.32, 0.80
CO_Y, CO_H = 4.28, 0.68
HAIR = 0.014


TITLE_H = 0.34


def node(name, cx, y, w, h, title, caption, *, accent=None, fill=WHITE, cap_color=SLATE):
    left = cx - w / 2
    c = deck.card(s, S, f"card{name}", left, y, w, h,
                  fill=fill, line=accent or MID, width_pt=1.75 if accent else 0.75)
    # Caption-less nodes centre their title; otherwise it sits above the caption.
    title_y = y + (0.16 if caption else (h - TITLE_H) / 2)
    ttf = deck.txbox(s, S, f"t{name}", left + 0.08, title_y, w - 0.16, TITLE_H, align=PP_ALIGN.CENTER)
    lines(ttf, title, align=PP_ALIGN.CENTER, font=HEAD, size=8.5, bold=True,
          color=accent or INK)
    if caption:
        ctf = deck.txbox(s, S, f"c{name}", left + 0.06, y + h - 0.28, w - 0.12, 0.22,
                         align=PP_ALIGN.CENTER)
        lines(ctf, caption, align=PP_ALIGN.CENTER, size=7, color=cap_color)
    return c


def vline(x, y0, y1, color=MID, weight=HAIR):
    deck.rule(s, x - weight / 2, y0, weight, y1 - y0, color=color)


def hline(x0, x1, y, color=MID, weight=HAIR):
    deck.rule(s, x0, y - weight / 2, x1 - x0, weight, color=color)


# ---------------------------------------------------------------- the tree
node("Root", MIDX, ROOT_Y, ROOT_W, ROOT_H, "Unicode Consortium", None)
vline(MIDX, ROOT_Y + ROOT_H, BUS1_Y)
hline(CENTERS[0], CENTERS[4], BUS1_Y)

COMMITTEES = [
    ("Board of\nDirectors", "Microsoft holds a seat", BLUE),
    ("CLDR Technical\nCommittee", "locale data", None),
    ("Unicode Technical\nCommittee", "Microsoft votes here", BLUE),
    ("ICU Technical\nCommittee", "ICU libraries", None),
    ("Editorial\nCommittee", "publications", None),
]
for i, (title, cap, accent) in enumerate(COMMITTEES):
    vline(CENTERS[i], BUS1_Y, ROW1_Y)
    node(f"Ch{i}", CENTERS[i], ROW1_Y, COLW, ROW1_H, title, cap,
         accent=accent, cap_color=BLUE if accent else SLATE)

# The working groups hang off the UTC alone, not off all five committees. The whole
# subtree is drawn in blue -- a thin grey bus here reads as a second fan-out from the
# Consortium and invites a column-wise "each committee owns a WG" misreading.
SPINE = 0.028
vline(MIDX, ROW1_Y + ROW1_H, BUS2_Y, color=BLUE, weight=SPINE)
hline(CENTERS[0], CENTERS[4], BUS2_Y, color=BLUE, weight=SPINE)

GROUPS = [
    ("Emoji Standard &\nResearch WG", "our proposal is read here", True),
    ("Script Encoding\nWG", "new scripts", False),
    ("CJK & Unihan\nWG", "Han characters", False),
    ("Properties &\nAlgorithms WG", "character properties", False),
    ("Editorial\nWG", "specification text", False),
]
for i, (title, cap, hot) in enumerate(GROUPS):
    vline(CENTERS[i], BUS2_Y, ROW2_Y, color=BLUE, weight=SPINE)
    node(f"Wg{i}", CENTERS[i], ROW2_Y, COLW, ROW2_H, title, cap,
         accent=AMBER if hot else None, cap_color=AMBER if hot else SLATE)

# "recommends" label, tucked into the gap between the UTC and its working-group bus
rtf = deck.txbox(s, S, "recl", MIDX + 0.12, ROW1_Y + ROW1_H + 0.005, 1.90, 0.15)
run(rtf.paragraphs[0], "working groups recommend; the UTC votes", size=7, italic=True, color=SLATE)

# ---------------------------------------------------------------- callout
deck.card(s, S, "cardCallout", 0.62, CO_Y, 8.76, CO_H, fill=WHITE, line=BLUE, width_pt=1.75)

t1 = deck.txbox(s, S, "co1", 0.82, CO_Y + 0.12, 8.36, 0.24)
run(t1.paragraphs[0], "Microsoft's vote sits one level above the working group that reads our proposal.",
    font=HEAD, size=10.5, bold=True, color=BLUE)
t2 = deck.txbox(s, S, "co2", 0.82, CO_Y + 0.40, 8.36, 0.22)
p2 = t2.paragraphs[0]
run(p2, "Full member = 1 UTC vote (Supporting = ½; associate, liaison and individual members do not vote)  ·  "
        "board seat: Vishal Chowdhary  ·  ships Segoe UI Emoji", size=8, color=SLATE)

deck.sources(s, S, "unicode.org/consortium/consort.html  ·  unicode.org/consortium/utc.html  ·  unicode.org/consortium/directors.html")

deck.save(OUT)
print(f"saved {OUT} | slides: {len(deck.slides)}")
print("\n--- geometry QA ---")
deck.qa(containers=("card",))
